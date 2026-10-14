"""Offline synthetic tests; no network, accounts or personal research files."""
import copy
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from weekly_common import load_config, load_json, dump_json, plain_text, expected_records, within
from vision_probe import make, grade, verify_capability
from update_weekly_docx import build
from validate_weekly_docx import validate_document
from docx import Document


class ShareTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=os.environ.get('WEEKLY_HIGHLIGHTS_TEST_TMP'))
        self.base = Path(self.temp.name)
        self.config = load_json(ROOT / 'config.example.json')
        self.config['workspace_root'] = str(self.base)
        self.config['execution'].update(host='synthetic-test', host_version='test', provider='test', review_model='test-vision', preferred_model='test-vision')
        self.config['setup']['status'] = 'vision_pending'
        self.config_path = self.base / 'config.json'
        dump_json(self.config_path, self.config)

    def tearDown(self):
        self.temp.cleanup()

    def response(self, challenge):
        # Unit testing the grader with explicit synthetic answers is NOT actual model validation.
        answers = load_json(self.base / 'probe' / 'answer.private.json')
        return {k: answers[k] for k in ['verification_code','leftmost_color','exponent','panels']} | {
            'actual_model':'test-vision', 'delivery':{'image_input_confirmed':True,'image_tool_evidence':['SYNTHETIC image1','SYNTHETIC image2']}}

    def test_unconfigured_or_untested_setup_blocks_runtime(self):
        with self.assertRaises((RuntimeError, ValueError)):
            load_config(self.config_path)
        self.assertEqual(load_config(self.config_path, require_ready=False)['setup']['status'], 'vision_pending')

    def test_self_claim_without_image_delivery_fails(self):
        challenge = make(self.config, self.base / 'probe')
        answer = self.response(challenge); answer['delivery'] = {'image_input_confirmed':False}
        dump_json(self.base/'response.json',answer)
        result = grade(self.base/'probe/challenge.json',self.base/'response.json',self.base/'receipt.json')
        self.assertEqual(result['status'], 'failed_or_unverified')

    def test_wrong_model_or_wrong_visual_answers_fail(self):
        challenge = make(self.config, self.base / 'probe')
        answer = self.response(challenge); answer['actual_model'] = 'different-model'
        dump_json(self.base/'response.json',answer)
        self.assertNotEqual(grade(self.base/'probe/challenge.json',self.base/'response.json',self.base/'receipt.json')['status'],'passed')

    def test_grader_pass_binds_model_and_images(self):
        challenge = make(self.config, self.base / 'probe')
        dump_json(self.base/'response.json',self.response(challenge))
        receipt = self.base/'receipt.json'
        self.assertEqual(grade(self.base/'probe/challenge.json',self.base/'response.json',receipt)['status'],'passed')
        self.config['setup'].update(status='ready',capability_receipt=str(receipt))
        self.assertEqual(verify_capability(self.config)['status'],'passed')
        self.config['execution']['review_model']='changed'
        with self.assertRaises(ValueError): verify_capability(self.config)

    def test_image_change_invalidates_capability(self):
        challenge = make(self.config, self.base/'probe')
        dump_json(self.base/'response.json',self.response(challenge))
        receipt = self.base/'receipt.json'; grade(self.base/'probe/challenge.json',self.base/'response.json',receipt)
        self.config['setup'].update(status='ready',capability_receipt=str(receipt))
        (self.base/'probe/input-check.png').write_bytes(b'changed')
        with self.assertRaises(ValueError): verify_capability(self.config)

    def test_new_empty_library_has_real_first_period(self):
        source = self.base/'JournalOne.docx'; doc=Document(); doc.add_paragraph('weekly highlights-Journal One'); doc.save(source)
        raw='SYNTHETIC TEST ONLY. Particle count 10<sup>6</sup> m<sup>-3</sup>. Not a scientific result.'
        records=[{'doi':'10.9999/test(only)','title':'SYNTHETIC SOFTWARE TEST','authors':['Test, Alice'],
                  'authors_verified':True,'authors_source_url':'https://example.org/test','abstract_html':raw,'abstract':plain_text(raw),'summary':'合成软件测试，不是真实论文。'}]
        metadata=self.base/'metadata.json';dump_json(metadata,records)
        output=self.base/'candidate.docx';build(source,expected_records(metadata),'2026.10.14',output)
        self.assertTrue(validate_document(output,metadata,'2026.10.14',source)['actual_fields_match'])
        self.assertEqual(len(Document(source).paragraphs),1)

    def test_paths_cannot_escape_selected_workspace(self):
        with self.assertRaises(ValueError):within(self.base,'../outside.docx')

    def test_scientific_text_is_preserved(self):
        self.assertEqual(plain_text('10<sup>6</sup> m<sup>-3</sup>, O<sub>3</sub>'),'10⁶ m⁻³, O₃')

if __name__=='__main__':unittest.main(verbosity=2)
