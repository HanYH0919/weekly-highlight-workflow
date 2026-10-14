"""Shared deterministic operations. No model calls or reference-manager access."""
import hashlib
import html
import json
import os
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

DOI_RE = re.compile(r'10\.\d{4,9}/[^\s<>"{}]+', re.I)
DATE_RE = re.compile(r'^20\d{2}\.\d{1,2}\.\d{1,2}$')
SUP = str.maketrans('0123456789+-=()abcdefghijklmnopqrstuvwxyz', '⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ᵃᵇᶜᵈᵉᶠᵍʰⁱʲᵏˡᵐⁿᵒᵖ𐞥ʳˢᵗᵘᵛʷˣʸᶻ')
SUB = str.maketrans('0123456789+-=()aehijklmnoprstuvx', '₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₕᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ')


def normalize_doi(value):
    value = html.unescape(str(value)).strip()
    if re.match(r'https?://(?:dx\.)?doi\.org/', value, re.I):
        value = unquote(re.sub(r'^https?://(?:dx\.)?doi\.org/', '', value, flags=re.I))
    match = DOI_RE.search(value)
    if not match:
        return ''
    result = match.group().rstrip('.,;')
    for left, right in [('(', ')'), ('[', ']')]:
        while result.endswith(right) and result.count(right) > result.count(left):
            result = result[:-1]
    return result.casefold()


def title_key(text):
    return ''.join(c for c in unicodedata.normalize('NFKC', plain_text(text)).casefold() if c.isalnum())


def scripted(text, kind):
    text = text.replace('−', '-').strip()
    mapping = SUP if kind == 'sup' else SUB
    if text and all(ord(c) in mapping for c in text):
        return text.translate(mapping)
    return ('^(' if kind == 'sup' else '_(') + text + ')'


def math_text(source):
    # Publisher fragments may omit the namespace declaration for a prefix.
    source = re.sub(r'(<\/?)[A-Za-z][\w-]*:', r'\1', source)
    root = ET.fromstring(source)

    def render(node):
        tag = node.tag.rsplit('}', 1)[-1]
        children = list(node)
        if tag in ('mi', 'mn', 'mo', 'mtext', 'ms'):
            return ''.join(node.itertext())
        if tag in ('math', 'mrow', 'mstyle', 'semantics', 'mpadded'):
            return (node.text or '') + ''.join(render(c) + (c.tail or '') for c in children if c.tag.rsplit('}', 1)[-1] not in ('annotation', 'annotation-xml'))
        if tag == 'msub' and len(children) == 2:
            return render(children[0]) + scripted(render(children[1]), 'sub')
        if tag == 'msup' and len(children) == 2:
            return render(children[0]) + scripted(render(children[1]), 'sup')
        if tag == 'msubsup' and len(children) == 3:
            return render(children[0]) + scripted(render(children[1]), 'sub') + scripted(render(children[2]), 'sup')
        if tag == 'mfrac' and len(children) == 2:
            return '(' + render(children[0]) + ')/(' + render(children[1]) + ')'
        if tag == 'msqrt':
            return '√(' + ''.join(render(c) for c in children) + ')'
        if tag in ('annotation', 'annotation-xml'):
            return ''
        if tag == 'mspace':
            return ' '
        raise ValueError(f'Unsupported MathML {tag}; verify original formula before continuing')

    return render(root)


class ScientificText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.scripts = []
        self.math = None
        self.math_tag = None

    def append(self, value):
        (self.scripts[-1][1] if self.scripts else self.parts).append(value)

    def handle_starttag(self, tag, attrs):
        if self.math is not None:
            self.math.append(self.get_starttag_text())
        elif tag.split(':')[-1] == 'math':
            self.math, self.math_tag = [self.get_starttag_text()], tag
        elif tag.split(':')[-1] in ('sub', 'sup'):
            self.scripts.append((tag.split(':')[-1], []))
        elif tag.split(':')[-1] in ('p', 'div', 'br', 'section'):
            self.append(' ')

    def handle_endtag(self, tag):
        if self.math is not None:
            self.math.append(f'</{tag}>')
            if tag == self.math_tag:
                source = ''.join(self.math)
                self.math = None
                self.append(math_text(source))
        elif tag.split(':')[-1] in ('sub', 'sup'):
            if not self.scripts or self.scripts[-1][0] != tag.split(':')[-1]:
                raise ValueError('Malformed scientific sub/sup markup')
            kind, pieces = self.scripts.pop()
            self.append(scripted(''.join(pieces), kind))
        elif tag.split(':')[-1] in ('p', 'div', 'section'):
            self.append(' ')

    def handle_startendtag(self, tag, attrs):
        if self.math is not None:
            self.math.append(self.get_starttag_text())
        else:
            self.handle_starttag(tag, attrs)

    def handle_data(self, value):
        if self.math is not None:
            self.math.append(html.escape(value))
        else:
            self.append(value)


def plain_text(value):
    value = html.unescape(str(value))
    # Only feed markup to HTMLParser: literal scientific comparisons remain literal.
    if re.search(r'<\/?[A-Za-z][\w:.-]*(?:\s[^>]*|/?)>', value):
        parser = ScientificText()
        parser.feed(value)
        parser.close()
        if parser.math is not None or parser.scripts:
            raise ValueError('Unclosed scientific markup')
        value = ''.join(parser.parts)
    return re.sub(r'\s+', ' ', html.unescape(value).replace('\xa0', ' ')).strip()


def format_authors(authors):
    if not authors:
        return ''
    if isinstance(authors, str):
        return plain_text(authors)
    names = []
    for author in authors:
        if isinstance(author, dict):
            name = ' '.join(filter(None, [author.get('given'), author.get('family')]))
        else:
            name = plain_text(author)
            if ',' in name:
                family, given = name.split(',', 1)
                name = given.strip() + ' ' + family.strip()
        names.append(name.strip())
    return names[0] if len(names) == 1 else ' and '.join(names) if len(names) == 2 else ', '.join(names[:-1]) + ', and ' + names[-1]


def load_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def dump_json(path, data, overwrite=False):
    path = Path(path)
    if path.exists() and not overwrite:
        raise FileExistsError(f'Refusing overwrite: {path}')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def digest(data):
    return hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode('utf-8')).hexdigest()


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def within(base, relative):
    base = Path(base).resolve()
    path = (base / relative).resolve()
    if not path.is_relative_to(base):
        raise ValueError(f'Path escapes intended directory: {path}')
    return path


def load_config(path=None, require_ready=True):
    if path:
        path = Path(path).expanduser()
    elif os.environ.get('WEEKLY_HIGHLIGHTS_CONFIG'):
        path = Path(os.environ['WEEKLY_HIGHLIGHTS_CONFIG']).expanduser()
    else:
        pointer = Path(__file__).resolve().parents[1] / 'installation.local.json'
        local = next((parent / '.weekly-highlights' / 'config.json' for parent in [Path.cwd(), *Path.cwd().parents]
                      if (parent / '.weekly-highlights' / 'config.json').is_file()), None)
        path = local or (Path(load_json(pointer)['config_path']) if pointer.is_file() else None)
    if path is None or not Path(path).is_file():
        raise RuntimeError('SETUP_REQUIRED: run scripts/setup_share.py configure; choose your workspace and models first')
    config = load_json(path)
    root = Path(config['workspace_root']).resolve()
    if not root.is_dir():
        raise FileNotFoundError(f'Workspace unavailable: {root}')
    if set(config['journals']) != {'ACP', 'PNAS'}:
        raise ValueError('Config must define ACP and PNAS')
    config['workspace_root'] = str(root)
    config['paths'] = {j: str(within(root, p)) for j, p in config['journals'].items()}
    config['_config_path'] = str(Path(path).resolve())
    if require_ready:
        from vision_probe import verify_capability
        if config.get('setup', {}).get('status') != 'ready':
            raise RuntimeError('VISION_CHECK_REQUIRED: finish the actual image test during setup; do not assume model vision from its name')
        verify_capability(config)
    return config


def date_window(config, today=None):
    try:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo(config['timezone'])
    except Exception:
        if config['timezone'] != 'Asia/Shanghai':
            raise
        tz = timezone(timedelta(hours=8))
    end = datetime.strptime(today, '%Y-%m-%d').date() if today else datetime.now(tz).date()
    return [(end - timedelta(days=int(config['window_days']) - 1)).isoformat(), end.isoformat()]


def extract_history(path):
    from docx import Document
    path = Path(path).resolve()
    before = file_hash(path)
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    rows, seen = [], set()
    for index, text in enumerate(paragraphs):
        if not re.match(r'^\s*(?:https?://(?:dx\.)?doi\.org/|DOI\s*[:：]?\s*|10\.\d{4,9}/)', text, re.I):
            continue
        doi = normalize_doi(text)
        if doi and doi not in seen:
            seen.add(doi)
            rows.append({'doi': doi, 'title': paragraphs[index - 2] if index >= 2 else '', 'authors': paragraphs[index - 1] if index >= 1 else '', 'source': str(path)})
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for match in DOI_RE.finditer(cell.text):
                    doi = normalize_doi(match.group())
                    if doi and doi not in seen:
                        seen.add(doi)
                        rows.append({'doi': doi, 'title': '', 'authors': '', 'source': str(path)})
    if before != file_hash(path):
        raise RuntimeError('Historical document changed while reading')
    return {'path': str(path), 'sha256': before, 'records': rows}


def expected_records(path):
    records = load_json(path)
    if not isinstance(records, list) or not records:
        raise ValueError('Metadata must be a non-empty list')
    records = [r for _, r in sorted(enumerate(records), key=lambda pair: int(pair[1].get('order', pair[0] + 1)))]
    output, seen = [], set()
    for record in records:
        if record.get('authors_verified') is not True or not record.get('authors_source_url'):
            raise ValueError('Word metadata requires verified authors and source')
        doi = normalize_doi(record.get('doi', ''))
        if not doi or doi in seen:
            raise ValueError('Invalid/duplicate metadata DOI')
        seen.add(doi)
        item = {k: plain_text(record.get(k, '')) for k in ('title', 'summary')}
        item['abstract'] = re.sub(r'^Abstract[.:]?\s+', '', plain_text(record.get('abstract_html') or record.get('abstract', '')), flags=re.I)
        if record.get('abstract_html') and record.get('abstract') and item['abstract'] != re.sub(r'^Abstract[.:]?\s+', '', plain_text(record['abstract']), flags=re.I):
            raise ValueError('Abstract text differs from preserved source markup')
        item.update(doi=doi, authors=format_authors(record.get('authors', [])))
        if any(not value for value in item.values()):
            raise ValueError('Empty Word field')
        output.append(item)
    return output


def pdf_relative(record):
    slug = re.sub(r'[<>:"/\\|?*]', '_', record['doi']).rstrip('. ')
    return record['type'] + '/' + slug + '.pdf'


def execution_records(approved, journal, evidence_path=None):
    """Allow evidence enrichment after review without changing confirmed scientific text."""
    records = [dict(r) for r in approved['articles'] if r['journal'] == journal]
    evidence = load_json(evidence_path) if evidence_path else []
    if isinstance(evidence, dict):
        evidence = evidence.get('articles', evidence.get('records', []))
    evidence = {normalize_doi(r.get('doi', '')): r for r in evidence}
    for record in records:
        extra = evidence.get(record['doi'], {})
        for key in ('title', 'summary', 'publication_date', 'journal', 'type'):
            if key in extra and extra[key] != record.get(key):
                raise ValueError(f'Execution evidence may not change approved {key}')
        if extra.get('abstract') and plain_text(extra['abstract']) != plain_text(record['abstract']):
            raise ValueError('Execution evidence may not change approved abstract')
        if extra.get('abstract_html') and plain_text(extra['abstract_html']) != plain_text(record['abstract']):
            raise ValueError('Scientific markup changes approved abstract; resolve and obtain revised approval')
        if record.get('authors') and extra.get('authors') and format_authors(record['authors']) != format_authors(extra['authors']):
            raise ValueError('Author conflict requires main-agent verification')
        for key in ('authors', 'authors_verified', 'authors_source_url', 'pdf_url', 'pdf_source_url', 'abstract_html'):
            if key in extra:
                record[key] = extra[key]
        record['pdf_file'] = record.get('pdf_file') or pdf_relative(record)
    return records
