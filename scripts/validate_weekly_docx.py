"""Validate actual five-field content, historical XML and numbering."""
import argparse
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from weekly_common import dump_json, expected_records


def validate_document(path, metadata, day, old_path=None):
    if isinstance(metadata, list):
        with tempfile.TemporaryDirectory() as temp:
            temp_path = Path(temp) / 'records.json'
            dump_json(temp_path, metadata)
            records = expected_records(temp_path)
    else:
        records = expected_records(metadata)
    doc = Document(str(path))
    paragraphs = doc.paragraphs
    if len(paragraphs) < 2 or paragraphs[1].text != day:
        raise AssertionError('Newest date must immediately follow document title')
    sequences = set()
    for index, record in enumerate(records):
        group = paragraphs[2 + 5 * index:7 + 5 * index]
        expected = [record['title'], record['authors'], 'https://doi.org/' + record['doi'],
                    'Summary：' + record['summary'], 'Abstract: ' + record['abstract']]
        if [p.text for p in group] != expected:
            raise AssertionError(f'Article {index + 1}: actual title/authors/DOI/Summary/Abstract differs from verified metadata')
        title, _, _, summary, abstract = group
        if not title.runs or title.runs[0].bold is not True:
            raise AssertionError('Title bold template lost')
        for paragraph, label in ((summary, 'Summary'), (abstract, 'Abstract')):
            if not paragraph.runs or paragraph.runs[0].text != label or paragraph.runs[0].bold is not True:
                raise AssertionError('Label format lost')
        if not title.paragraph_format.keep_together:
            raise AssertionError('Multi-line title can split across pages')
        sequences.add(title._p.pPr.numPr.numId.val)
    if len(sequences) != 1:
        raise AssertionError('New articles must share one independent numbering instance')
    new_id = next(iter(sequences))
    num = next(n for n in doc.part.numbering_part.element.findall(qn('w:num')) if int(n.get(qn('w:numId'))) == new_id)
    starts = num.findall('.//' + qn('w:startOverride'))
    if not starts or starts[0].get(qn('w:val')) != '1':
        raise AssertionError('New date numbering must start at 1')
    if old_path:
        old = Document(str(old_path))
        tail = paragraphs[2 + len(records) * 5:]
        canonical = lambda p: etree.tostring(p._p, method='c14n')
        if [canonical(p) for p in tail] != [canonical(p) for p in old.paragraphs[1:]]:
            raise AssertionError('Historical paragraph content or formatting changed')
        if canonical(paragraphs[0]) != canonical(old.paragraphs[0]):
            raise AssertionError('Document title changed')
        if len(old.tables) != len(doc.tables) or [etree.tostring(t._tbl, method='c14n') for t in old.tables] != [etree.tostring(t._tbl, method='c14n') for t in doc.tables]:
            raise AssertionError('Historical tables changed')
        old_num_ids = {int(n.get(qn('w:numId'))) for n in old.part.numbering_part.element.findall(qn('w:num'))}
        if new_id in old_num_ids:
            raise AssertionError('New date reuses historical numbering')
    return {'articles': len(records), 'actual_fields_match': True, 'old_content_and_format_preserved': bool(old_path), 'independent_numbering': True}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--docx', type=Path, required=True)
    p.add_argument('--metadata', type=Path, required=True)
    p.add_argument('--date', required=True)
    p.add_argument('--old-docx', type=Path)
    args = p.parse_args()
    print('VALIDATION_PASS', validate_document(args.docx, args.metadata, args.date, args.old_docx))

if __name__ == '__main__':
    main()
